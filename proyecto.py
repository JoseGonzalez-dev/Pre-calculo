"""
Proyecto Final de Precálculo: Escalamiento y Servidores de un Sistema en la Nube
------------------------------------------------------------------------------
Modelos:
    U(t) = U0 * (1 + r)^t          -> crecimiento de usuarios (exponencial)
    L(u) = 100 / (CAP - u)         -> latencia en ms (racional, asíntota en u = CAP)
    P(x) = -0.5x^2 + 1000x - 5000  -> utilidad financiera (cuadrática)

Uso:
    pip install numpy matplotlib
    python proyecto_precalculo.py            # consola + gráficas interactivas
    python proyecto_precalculo.py --excel    # genera tablas_proyecto.xlsx (necesita openpyxl)

Muestra el desarrollo paso a paso en consola y abre una ventana con las 3 gráficas
y sliders para cambiar la tasa de crecimiento, los usuarios iniciales y la capacidad.
"""

import math
import sys

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.widgets import Slider

# ---------------- Parámetros base (los del enunciado) ----------------
U0 = 150          # usuarios iniciales
R = 0.25          # crecimiento mensual (25 %)
META = 3000       # usuarios de la Pregunta 1
CAP = 5000        # capacidad (asíntota vertical)
NUM = 100         # numerador de la latencia
LAT_MAX_MS = 2000 # 2 segundos, Pregunta 2
A, B, C = -0.5, 1000, -5000  # coeficientes de P(x)


# ---------------- Modelos ----------------
def usuarios(t, u0=U0, r=R):
    return u0 * (1 + r) ** t


def latencia(u, cap=CAP):
    return NUM / (cap - u)


def utilidad(x):
    return A * x**2 + B * x + C


# ---------------- Soluciones analíticas ----------------
def mes_para_meta(u0=U0, r=R, meta=META):
    """Despeja t: meta = u0*(1+r)^t  ->  t = ln(meta/u0) / ln(1+r)"""
    return math.log(meta / u0) / math.log(1 + r)


def usuarios_limite_latencia(cap=CAP, lat_ms=LAT_MAX_MS):
    """L(u) > lat  ->  NUM/(cap-u) > lat  ->  u > cap - NUM/lat  (con u < cap)"""
    return cap - NUM / lat_ms


def vertice():
    h = -B / (2 * A)
    k = utilidad(h)
    return h, k


def mes_colapso(u0=U0, r=R, cap=CAP):
    """Mes en el que U(t) alcanza la asíntota."""
    return math.log(cap / u0) / math.log(1 + r)


# ---------------- Desarrollo paso a paso en consola ----------------
def imprimir_desarrollo():
    t = mes_para_meta()
    print("=" * 60)
    print("PREGUNTA 1: ¿En qué mes se superan 3,000 usuarios?")
    print("  150 * (1.25)^t = 3000")
    print("  (1.25)^t = 3000 / 150 = 20")
    print("  t * ln(1.25) = ln(20)")
    print(f"  t = ln(20) / ln(1.25) = {math.log(20):.4f} / {math.log(1.25):.4f} = {t:.4f}")
    print(f"  -> Se supera durante el mes {math.ceil(t)}")
    print(f"     U(13) = {usuarios(13):,.1f}   U(14) = {usuarios(14):,.1f}")

    u = usuarios_limite_latencia()
    print("\n" + "=" * 60)
    print("PREGUNTA 2: ¿Desde cuántos usuarios la latencia supera 2000 ms?")
    print("  100 / (5000 - u) > 2000")
    print("  100 > 2000 * (5000 - u)      (5000 - u > 0, no cambia el sentido)")
    print("  0.05 > 5000 - u")
    print(f"  u > 5000 - 0.05 = {u:.2f}")
    print("  -> Con u < 5000: la latencia supera 2 s entre 4,999.95 y 5,000 usuarios.")

    h, k = vertice()
    print("\n" + "=" * 60)
    print("PREGUNTA 3: Precio que maximiza la utilidad")
    print("  P(x) = -0.5x^2 + 1000x - 5000   (a < 0: abre hacia abajo, hay máximo)")
    print("  h = -b / (2a) = -1000 / (2 * -0.5) = 1000")
    print(f"  k = P(1000) = -0.5(1000)^2 + 1000(1000) - 5000 = {k:,.0f}")
    print(f"  -> Precio óptimo: ${h:,.0f}   Utilidad máxima: ${k:,.0f}")

    tc = mes_colapso()
    print("\n" + "=" * 60)
    print("PARA LAS CONCLUSIONES")
    print(f"  U(t) llega a 5,000 usuarios en t = {tc:.2f} meses (mes {math.ceil(tc)}).")
    print(f"  U(15) = {usuarios(15):,.0f} usuarios -> hay que escalar antes del mes 15.")
    print("=" * 60)


# ---------------- Exportar tablas a Excel ----------------
def exportar_excel(ruta="tablas_proyecto.xlsx"):
    """Genera un .xlsx con parámetros, resultados y tablas de valores.
    Todo usa fórmulas ligadas a la hoja 'Parámetros': si cambias un valor ahí,
    las tablas y los resultados se recalculan solos."""
    from openpyxl import Workbook
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.styles import Alignment, Font, PatternFill

    f_base = Font(name="Arial", size=10)
    f_head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    f_input = Font(name="Arial", size=10, color="0000FF")
    f_title = Font(name="Arial", size=13, bold=True)
    fill_head = PatternFill("solid", start_color="1F4E78")

    wb = Workbook()
    ws_p = wb.active
    ws_p.title = "Parámetros"
    ws_r = wb.create_sheet("Resultados")
    ws_u = wb.create_sheet("Usuarios")
    ws_l = wb.create_sheet("Latencia")
    ws_g = wb.create_sheet("Utilidad")

    def encabezado(ws, fila, textos):
        for i, txt in enumerate(textos, start=1):
            c = ws.cell(row=fila, column=i, value=txt)
            c.font, c.fill = f_head, fill_head
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # --- Parámetros (celdas azules = se pueden editar) ---
    ws_p["A1"] = "Parámetros del modelo (editar solo las celdas azules)"
    ws_p["A1"].font = f_title
    encabezado(ws_p, 3, ["Parámetro", "Valor", "Descripción"])
    params = [
        ("U0", U0, "Usuarios iniciales"),
        ("r", R, "Crecimiento mensual (0.25 = 25 %)"),
        ("Meta de usuarios", META, "Pregunta 1"),
        ("Capacidad", CAP, "Asíntota vertical de L(u): colapso del servidor"),
        ("Numerador de L(u)", NUM, "L(u) = numerador / (capacidad - u), en ms"),
        ("Latencia máxima (ms)", LAT_MAX_MS, "Pregunta 2: 2 s = 2000 ms"),
        ("a", A, "Coeficiente de x² en P(x)"),
        ("b", B, "Coeficiente de x en P(x)"),
        ("c", C, "Término independiente de P(x)"),
    ]
    for i, (nombre, valor, desc) in enumerate(params, start=4):
        ws_p.cell(row=i, column=1, value=nombre).font = f_base
        c = ws_p.cell(row=i, column=2, value=valor)
        c.font = f_input
        ws_p.cell(row=i, column=3, value=desc).font = f_base
    ws_p["A14"] = "Fuente: valores del enunciado 'Proyecto Final: Escalamiento y Servidores de un Sistema en la Nube'."
    ws_p["A14"].font = Font(name="Arial", size=9, italic=True)

    p = {k: f"'Parámetros'!$B${i}" for k, i in
         dict(u0=4, r=5, meta=6, cap=7, num=8, lat=9, a=10, b=11, c=12).items()}

    # --- Resultados ---
    ws_r["A1"] = "Resultados de las 3 preguntas"
    ws_r["A1"].font = f_title
    encabezado(ws_r, 3, ["Resultado", "Valor", "Fórmula"])
    res = [
        ("P1: mes exacto en que U(t) = meta", f"=LN({p['meta']}/{p['u0']})/LN(1+{p['r']})",
         "t = ln(meta/U0) / ln(1+r)", "0.0000"),
        ("P1: primer mes entero que supera la meta", "=INT(B4)+1",
         "Se redondea hacia arriba el mes exacto", "0"),
        ("P2: usuarios a partir de los cuales L(u) > 2000 ms", f"={p['cap']}-{p['num']}/{p['lat']}",
         "u > capacidad - numerador / latencia", "#,##0.00"),
        ("P3: precio óptimo h ($)", f"=-{p['b']}/(2*{p['a']})", "h = -b / (2a)", "$#,##0"),
        ("P3: utilidad máxima k ($)", f"={p['a']}*B7^2+{p['b']}*B7+{p['c']}", "k = P(h)", "$#,##0"),
        ("Mes en que U(t) llega a la capacidad", f"=LN({p['cap']}/{p['u0']})/LN(1+{p['r']})",
         "t = ln(capacidad/U0) / ln(1+r)", "0.0000"),
        ("Usuarios en el último mes completo antes del colapso",
         f"={p['u0']}*(1+{p['r']})^INT(B9)", "U(INT(t_colapso))", "#,##0"),
    ]
    for i, (nombre, formula, nota, fmt) in enumerate(res, start=4):
        ws_r.cell(row=i, column=1, value=nombre).font = f_base
        c = ws_r.cell(row=i, column=2, value=formula)
        c.font, c.number_format = f_base, fmt
        ws_r.cell(row=i, column=3, value=nota).font = f_base

    # --- Usuarios: U(t) ---
    encabezado(ws_u, 1, ["t (meses)", "U(t) usuarios", "¿Supera la meta?", "¿Supera la capacidad?"])
    for t in range(0, 25):
        fila = t + 2
        ws_u.cell(row=fila, column=1, value=t).font = f_base
        c = ws_u.cell(row=fila, column=2, value=f"={p['u0']}*(1+{p['r']})^A{fila}")
        c.font, c.number_format = f_base, "#,##0.0"
        ws_u.cell(row=fila, column=3, value=f'=IF(B{fila}>{p["meta"]},"Sí","No")').font = f_base
        ws_u.cell(row=fila, column=4, value=f'=IF(B{fila}>{p["cap"]},"Sí","No")').font = f_base

    # --- Latencia: L(u) ---
    encabezado(ws_l, 1, ["u (usuarios)", "L(u) en ms", "¿Supera 2000 ms?"])
    usuarios_tabla = [0, 1000, 2000, 3000, 4000, 4500, 4900, 4990, 4999, 4999.9,
                      4999.95, 4999.96, 4999.99, 5000]
    for i, u in enumerate(usuarios_tabla, start=2):
        ws_l.cell(row=i, column=1, value=u).font = f_base
        c = ws_l.cell(row=i, column=2,
                      value=f'=IF(A{i}<{p["cap"]},{p["num"]}/({p["cap"]}-A{i}),"Indefinida")')
        c.font, c.number_format = f_base, "#,##0.0000"
        ws_l.cell(row=i, column=3,
                  value=f'=IF(ISNUMBER(B{i}),IF(B{i}>{p["lat"]},"Sí","No"),"Colapso")').font = f_base
    nota = len(usuarios_tabla) + 3
    ws_l.cell(row=nota, column=1,
              value="En u = 4999.95 la latencia es exactamente 2000 ms (no la supera). "
                    "u = 5000 es la asíntota: la función no está definida.").font = \
        Font(name="Arial", size=9, italic=True)

    # --- Utilidad: P(x) ---
    encabezado(ws_g, 1, ["x (precio, $)", "P(x) utilidad ($)"])
    for i, x in enumerate(range(0, 2001, 100), start=2):
        ws_g.cell(row=i, column=1, value=x).font = f_base
        c = ws_g.cell(row=i, column=2, value=f"={p['a']}*A{i}^2+{p['b']}*A{i}+{p['c']}")
        c.font, c.number_format = f_base, "$#,##0;($#,##0);-"

    # --- Gráficas dentro de Excel ---
    def grafica(ws, titulo, tx, ty, ult_fila, ancla):
        ch = ScatterChart()
        ch.title, ch.style = titulo, 13
        ch.x_axis.title, ch.y_axis.title = tx, ty
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        xs = Reference(ws, min_col=1, min_row=2, max_row=ult_fila)
        ys = Reference(ws, min_col=2, min_row=1, max_row=ult_fila)
        s = Series(ys, xs, title_from_data=True)
        s.marker.symbol = "none"
        ch.series.append(s)
        ch.height, ch.width = 9, 16
        ws.add_chart(ch, ancla)

    grafica(ws_u, "Crecimiento de usuarios U(t)", "t (meses)", "Usuarios", 26, "F2")
    grafica(ws_g, "Utilidad financiera P(x)", "Precio x ($)", "Utilidad ($)", 22, "D2")

    # --- Anchos de columna ---
    for ws, anchos in ((ws_p, (24, 12, 52)), (ws_r, (52, 16, 42)), (ws_u, (12, 16, 18, 22)),
                       (ws_l, (16, 18, 20)), (ws_g, (16, 20))):
        for i, w in enumerate(anchos, start=1):
            ws.column_dimensions[chr(64 + i)].width = w

    wb.save(ruta)
    print(f"Excel generado: {ruta}")


# ---------------- Gráficas interactivas ----------------
def main():
    if "--excel" in sys.argv:
        exportar_excel()
        return
    imprimir_desarrollo()

    fig, (ax1, ax2, ax3) = plt.subplots(1, 3, figsize=(16, 7.5))
    fig.subplots_adjust(bottom=0.36, wspace=0.28, top=0.90)
    fig.suptitle("Escalamiento y servidores de un sistema en la nube", fontsize=14)

    def nuevo_slider(y, etiqueta, vmin, vmax, vinit, paso):
        ax = fig.add_axes([0.17, y, 0.68, 0.025])
        return Slider(ax, etiqueta, vmin, vmax, valinit=vinit, valstep=paso)

    s_r = nuevo_slider(0.25, "Crecimiento mensual", 0.05, 0.60, R, 0.01)
    s_u0 = nuevo_slider(0.20, "Usuarios iniciales", 50, 1000, U0, 10)
    s_cap = nuevo_slider(0.15, "Capacidad (asíntota)", 2000, 10000, CAP, 100)
    s_t = nuevo_slider(0.10, "Mes t (punto en 1 y 2)", 0, 24, 12, 0.1)
    s_x = nuevo_slider(0.05, "Precio x (punto en 3)", 0, 2000, 500, 10)

    def dibujar(_=None):
        r, u0, cap = s_r.val, s_u0.val, s_cap.val
        t_sel, x_sel = s_t.val, s_x.val
        u_sel = usuarios(t_sel, u0, r)

        # ---- 1) Exponencial ----
        ax1.clear()
        tmax = max(mes_colapso(u0, r, cap) * 1.3, mes_para_meta(u0, r) * 1.3, 6) if u0 < cap else 24
        tmax = max(tmax, 24)
        t = np.linspace(0, tmax, 400)
        ax1.plot(t, usuarios(t, u0, r), color="tab:blue", label="U(t)")
        ax1.axhline(META, color="green", ls="--", lw=1, label=f"Meta {META:,}")
        ax1.axhline(cap, color="red", ls="--", lw=1, label=f"Capacidad {cap:,.0f}")
        if u0 < META:
            tm = mes_para_meta(u0, r)
            ax1.plot(tm, META, "go")
            ax1.annotate(f"t = {tm:.2f}\n(mes {math.ceil(tm)})", (tm, META),
                         textcoords="offset points", xytext=(-70, 12))
        if u0 < cap:
            tc = mes_colapso(u0, r, cap)
            ax1.plot(tc, cap, "ro")
            ax1.annotate(f"colapso\nt = {tc:.2f}", (tc, cap),
                         textcoords="offset points", xytext=(-60, 12), color="red")
        ax1.axvline(t_sel, color="black", ls=":", lw=1)
        ax1.plot(t_sel, u_sel, "ko")
        ax1.set_ylim(0, cap * 1.4)
        ax1.set_title("1) Exponencial: crecimiento de usuarios")
        ax1.set_xlabel(f"Tiempo t (meses)   |   t = {t_sel:.1f} -> U = {u_sel:,.0f} usuarios")
        ax1.set_ylabel("Usuarios")
        ax1.legend(loc="upper left", fontsize=8)
        ax1.grid(alpha=0.3)

        # ---- 2) Racional ----
        ax2.clear()
        u = cap - np.logspace(np.log10(cap), -3, 800)  # 0 -> cap - 0.001
        ax2.semilogy(u, latencia(u, cap), color="tab:orange", label="L(u)")
        ax2.axvline(cap, color="red", ls="--", lw=1, label=f"Asíntota u = {cap:,.0f}")
        ax2.axhline(LAT_MAX_MS, color="purple", ls="--", lw=1, label="2000 ms (2 s)")
        ul = usuarios_limite_latencia(cap)
        ax2.plot(ul, LAT_MAX_MS, "mo")
        ax2.annotate(f"u = {ul:,.2f}", (ul, LAT_MAX_MS),
                     textcoords="offset points", xytext=(-90, 10), color="purple")
        if u_sel < cap:
            lat = latencia(u_sel, cap)
            ax2.axvline(u_sel, color="black", ls=":", lw=1)
            ax2.plot(u_sel, lat, "ko")
            estado = f"L = {lat:,.4f} ms" + ("  (> 2 s)" if lat > LAT_MAX_MS else "")
        else:
            estado = "COLAPSO: usuarios >= capacidad"
        ax2.text(0.03, 0.74, f"Mes t = {t_sel:.1f}\nUsuarios = {u_sel:,.0f}\n{estado}",
                 transform=ax2.transAxes, fontsize=9, va="top",
                 color="red" if u_sel >= cap else "black",
                 bbox=dict(boxstyle="round", fc="white", ec="gray", alpha=0.9))
        ax2.set_ylim(1e-2, 1e4)
        ax2.set_title("2) Racional: latencia del servidor (escala log)")
        ax2.set_xlabel("Usuarios activos u")
        ax2.set_ylabel("Latencia (ms)")
        ax2.legend(loc="upper left", fontsize=8)
        ax2.grid(alpha=0.3, which="both")

        # ---- 3) Cuadrática ----
        ax3.clear()
        x = np.linspace(0, 2000, 400)
        ax3.plot(x, utilidad(x), color="tab:green", label="P(x)")
        h, k = vertice()
        ax3.plot(h, k, "ro")
        ax3.axvline(h, color="gray", ls=":", lw=1)
        ax3.annotate(f"Vértice ({h:,.0f}, {k:,.0f})", (h, k),
                     textcoords="offset points", xytext=(-70, 12), color="red")
        ax3.plot(x_sel, utilidad(x_sel), "ko", label="Precio elegido")
        ax3.axhline(0, color="black", lw=0.8)
        ax3.set_ylim(-20000, 560000)
        ax3.set_title("3) Cuadrática: utilidad financiera")
        ax3.set_xlabel(f"Precio x ($)   |   x = ${x_sel:,.0f} -> P = ${utilidad(x_sel):,.0f}")
        ax3.set_ylabel("Utilidad P(x) ($)")
        ax3.legend(loc="lower center", fontsize=8)
        ax3.grid(alpha=0.3)

        fig.canvas.draw_idle()

    for s in (s_r, s_u0, s_cap, s_t, s_x):
        s.on_changed(dibujar)

    dibujar()
    plt.show()


if __name__ == "__main__":
    main()
